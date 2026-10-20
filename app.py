import csv, io, json, os, re, secrets, sqlite3, hashlib, hmac, time, zipfile, math
from pathlib import Path
from contextlib import contextmanager,asynccontextmanager
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from urllib.parse import urlparse
from fastapi import FastAPI, Request, HTTPException, UploadFile, File, Form, Depends
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
import httpx
from review_pipeline import prepare_document, normalize_formal, normalize_score, return_reason, FORMAL_PROMPT, SCORING_PROMPT, DEFAULT_STANDARD
from paper_queue import ReviewQueue, BUSY, BJT, next_daily, peak, DeferReview
import writing_workflow

ROOT=Path(__file__).parent
DATA=Path(os.getenv('DATA_DIR', str(ROOT/'data')))
DATA.mkdir(parents=True,exist_ok=True)
(DATA/'files').mkdir(exist_ok=True)
os.chmod(DATA,0o700)
DB=DATA/'course.db'
LIMIT=5_000_000
SECURE=os.getenv('COOKIE_SECURE','false').lower()=='true'
POOL=ThreadPoolExecutor(max_workers=1)
@asynccontextmanager
async def lifespan(application):
    REVIEW_QUEUE.start()
    try:yield
    finally:REVIEW_QUEUE.stop()
app=FastAPI(title='智学工坊 · 课程教学平台',docs_url=None,redoc_url=None,lifespan=lifespan)

@contextmanager
def db():
    c=sqlite3.connect(DB,timeout=20)
    c.row_factory=sqlite3.Row
    c.execute('PRAGMA foreign_keys=ON')
    try:
        yield c
        c.commit()
    except Exception:
        c.rollback(); raise
    finally: c.close()

def now(): return datetime.now().isoformat(timespec='seconds')
def pwdhash(p):
    salt=secrets.token_hex(16)
    return salt+':'+hashlib.scrypt(p.encode(),salt=bytes.fromhex(salt),n=16384,r=8,p=1).hex()
def pwdcheck(p,stored):
    s,v=stored.split(':')
    return hmac.compare_digest(v,hashlib.scrypt(p.encode(),salt=bytes.fromhex(s),n=16384,r=8,p=1).hex())
def init():
    with db() as c:
        c.executescript('''
        PRAGMA journal_mode=WAL;
        CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY,name TEXT NOT NULL,student_no TEXT UNIQUE NOT NULL,password TEXT NOT NULL,role TEXT NOT NULL DEFAULT 'student',created TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS sessions(token TEXT PRIMARY KEY,user_id INTEGER REFERENCES users(id),expires REAL);
        CREATE TABLE IF NOT EXISTS chapters(id INTEGER PRIMARY KEY,title TEXT,description TEXT,objective TEXT);
        CREATE TABLE IF NOT EXISTS materials(id INTEGER PRIMARY KEY,chapter_id INTEGER REFERENCES chapters(id),title TEXT,path TEXT,filename TEXT,created TEXT);
        CREATE TABLE IF NOT EXISTS questions(id INTEGER PRIMARY KEY,chapter_id INTEGER REFERENCES chapters(id),level INTEGER,stem TEXT,options TEXT,answer INTEGER,explanation TEXT);
        CREATE TABLE IF NOT EXISTS attempts(id INTEGER PRIMARY KEY,user_id INTEGER REFERENCES users(id),question_id INTEGER REFERENCES questions(id),choice INTEGER,score REAL,created TEXT);
        CREATE TABLE IF NOT EXISTS assignments(id INTEGER PRIMARY KEY,title TEXT,instructions TEXT,rubric TEXT,deadline TEXT,created TEXT);
        CREATE TABLE IF NOT EXISTS submissions(id INTEGER PRIMARY KEY,assignment_id INTEGER REFERENCES assignments(id),user_id INTEGER REFERENCES users(id),path TEXT,filename TEXT,created TEXT,status TEXT DEFAULT '待批改',ai_json TEXT,score REAL,feedback TEXT DEFAULT '',published INTEGER DEFAULT 0, UNIQUE(assignment_id,user_id));
        CREATE TABLE IF NOT EXISTS throttle(bucket TEXT,stamp REAL);
        CREATE INDEX IF NOT EXISTS throttle_idx ON throttle(bucket,stamp);
        CREATE INDEX IF NOT EXISTS attempts_idx ON attempts(user_id,question_id);
        CREATE TABLE IF NOT EXISTS audit(id INTEGER PRIMARY KEY,user_id INTEGER,action TEXT,target TEXT,created TEXT);
        CREATE TABLE IF NOT EXISTS ai_practice(id INTEGER PRIMARY KEY,user_id INTEGER,chapter_id INTEGER,payload TEXT,choice INTEGER,created TEXT);
        ''')
        c.execute("CREATE TABLE IF NOT EXISTS learner_profiles(user_id INTEGER PRIMARY KEY REFERENCES users(id),research_area TEXT DEFAULT '',writing_stage TEXT DEFAULT '选题阶段',goal TEXT DEFAULT '',teacher_note TEXT DEFAULT '',updated TEXT)")
        c.execute("CREATE TABLE IF NOT EXISTS submission_history(id INTEGER PRIMARY KEY,submission_id INTEGER,user_id INTEGER,path TEXT,filename TEXT,snapshot TEXT,archived TEXT)")
        assignment_columns={row['name'] for row in c.execute('PRAGMA table_info(assignments)')}
        if 'grading_instructions' not in assignment_columns:c.execute("ALTER TABLE assignments ADD COLUMN grading_instructions TEXT NOT NULL DEFAULT ''")
        if 'revision' not in assignment_columns:c.execute('ALTER TABLE assignments ADD COLUMN revision INTEGER NOT NULL DEFAULT 1')
        if 'workflow' not in assignment_columns:c.execute('ALTER TABLE assignments ADD COLUMN workflow INTEGER NOT NULL DEFAULT 0')
        writing_workflow.migrate(c)
        columns={row['name'] for row in c.execute('PRAGMA table_info(submissions)')}
        if 'returned' not in columns:c.execute('ALTER TABLE submissions ADD COLUMN returned INTEGER NOT NULL DEFAULT 0')
        if 'return_reason' not in columns:c.execute("ALTER TABLE submissions ADD COLUMN return_reason TEXT NOT NULL DEFAULT ''")
        if 'returned_at' not in columns:c.execute("ALTER TABLE submissions ADD COLUMN returned_at TEXT NOT NULL DEFAULT ''")
        c.execute("INSERT OR IGNORE INTO learner_profiles(user_id,updated) SELECT id,? FROM users WHERE role='student'",(now(),))
        if 'appeal_reason' not in columns:c.execute("ALTER TABLE submissions ADD COLUMN appeal_reason TEXT NOT NULL DEFAULT ''")
        c.execute("CREATE TABLE IF NOT EXISTS ai_usage(id INTEGER PRIMARY KEY,user_id INTEGER,submission_id INTEGER,purpose TEXT,model TEXT,input_tokens INTEGER,output_tokens INTEGER,cache_hit_tokens INTEGER,local_cache INTEGER,created TEXT)")
        c.execute("CREATE TABLE IF NOT EXISTS ai_result_cache(hash TEXT PRIMARY KEY,payload TEXT,created TEXT)")
        if not c.execute('SELECT 1 FROM chapters').fetchone():
            from seed import CHAPTERS,QUESTIONS
            c.executemany('INSERT INTO chapters VALUES(?,?,?,?)',CHAPTERS)
            c.executemany('INSERT INTO questions(chapter_id,level,stem,options,answer,explanation) VALUES(?,?,?,?,?,?)',[(ch,lv,s,json.dumps(op,ensure_ascii=False),a,e) for ch,lv,s,op,a,e in QUESTIONS])
            rubric=json.dumps([{'name':'研究问题与目标','max':20},{'name':'结构与逻辑','max':25},{'name':'证据与论证','max':25},{'name':'学术表达','max':20},{'name':'引用规范','max':10}],ensure_ascii=False)
            c.execute('INSERT INTO assignments(title,instructions,rubric,deadline,created) VALUES(?,?,?,?,?)',('论文写作综合作业','请选择与你研究方向相关的主题，提交论文草稿或完整的引言与方法部分。明确研究问题、研究空白、方法和证据；参考文献须真实、可核查。支持 DOCX、文字型 PDF，不超过5 MB。',rubric,'',now()))
        c.execute("UPDATE submissions SET status='AI任务中断，可重新批改' WHERE status IN ('AI批改中','AI排队中')")
init()

@app.middleware('http')
async def guard(request,call_next):
    if request.method not in ('GET','HEAD','OPTIONS') and request.url.path.startswith('/api/'):
        if request.headers.get('X-App-Request')!='1':return JSONResponse({'detail':'请求校验失败，请刷新页面'},403)
        origin=request.headers.get('origin')
        if origin and urlparse(origin).netloc!=request.headers.get('host'):return JSONResponse({'detail':'不允许跨站请求'},403)
    result=await call_next(request)
    result.headers['X-Content-Type-Options']='nosniff'
    result.headers['Referrer-Policy']='same-origin'
    result.headers['X-Frame-Options']='SAMEORIGIN'
    if request.url.path.startswith('/api/'): result.headers['Cache-Control']='no-store'
    return result

def throttle(bucket,limit,window=600):
    with db() as c:
        c.execute('DELETE FROM throttle WHERE stamp<?',(time.time()-86400,))
        n=c.execute('SELECT COUNT(*) FROM throttle WHERE bucket=? AND stamp>?',(bucket,time.time()-window)).fetchone()[0]
        if n>=limit:raise HTTPException(429,'操作过于频繁，请稍后再试')
        c.execute('INSERT INTO throttle VALUES(?,?)',(bucket,time.time()))

def current(request:Request):
    token=request.cookies.get('session','')
    with db() as c:
        u=c.execute('SELECT u.id,u.name,u.student_no,u.role FROM users u JOIN sessions s ON s.user_id=u.id WHERE s.token=? AND s.expires>?',(hashlib.sha256(token.encode()).hexdigest(),time.time())).fetchone()
    if not u:raise HTTPException(401,'请先登录')
    return dict(u)
def teacher(u=Depends(current)):
    if u['role']!='teacher':raise HTTPException(403,'仅教师可操作')
    return u

def audit(uid,action,target):
    with db() as c:c.execute('INSERT INTO audit(user_id,action,target,created) VALUES(?,?,?,?)',(uid,action,str(target),now()))

class Register(BaseModel):
    name:str=Field(min_length=2,max_length=40)
    student_no:str=Field(min_length=4,max_length=30)
    password:str=Field(min_length=10,max_length=128)
class Login(BaseModel):
    student_no:str=Field(min_length=1,max_length=30)
    password:str=Field(min_length=1,max_length=128)
@app.post('/api/register')
def register(x:Register,request:Request):
    throttle('register:'+request.client.host,20,3600)
    if not re.fullmatch(r'[A-Za-z0-9_-]+',x.student_no):raise HTTPException(400,'学号仅支持字母、数字、下划线和短横线')
    if not x.name.strip():raise HTTPException(400,'请填写姓名')
    try:
        with db() as c:
            uid=c.execute('INSERT INTO users(name,student_no,password,created) VALUES(?,?,?,?)',(x.name.strip(),x.student_no,pwdhash(x.password),now())).lastrowid
            c.execute('INSERT INTO learner_profiles(user_id,updated) VALUES(?,?)',(uid,now()))
    except sqlite3.IntegrityError:raise HTTPException(409,'学号已注册，请登录或联系教师重置密码')
    return {'ok':True}
@app.post('/api/login')
def login(x:Login,request:Request):
    throttle('login-ip:'+request.client.host,50)
    throttle('login-account:'+x.student_no,15)
    with db() as c:
        u=c.execute('SELECT * FROM users WHERE student_no=?',(x.student_no,)).fetchone()
        if not u or not pwdcheck(x.password,u['password']):raise HTTPException(401,'账号或密码不正确')
        token=secrets.token_urlsafe(40)
        c.execute('DELETE FROM sessions WHERE expires<?',(time.time(),))
        c.execute('INSERT INTO sessions VALUES(?,?,?)',(hashlib.sha256(token.encode()).hexdigest(),u['id'],time.time()+43200))
    response=JSONResponse({'ok':True})
    response.set_cookie('session',token,httponly=True,secure=SECURE,samesite='lax',max_age=43200)
    return response
@app.post('/api/logout')
def logout(request:Request):
    with db() as c:c.execute('DELETE FROM sessions WHERE token=?',(hashlib.sha256(request.cookies.get('session','').encode()).hexdigest(),))
    r=JSONResponse({'ok':True});r.delete_cookie('session');return r
@app.get('/api/me')
def me(u=Depends(current)):return {**u,'secure':SECURE}
class Password(BaseModel):
    old:str=Field(max_length=128)
    new:str=Field(min_length=10,max_length=128)
@app.post('/api/password')
def password(x:Password,u=Depends(current)):
    with db() as c:
        old=c.execute('SELECT password FROM users WHERE id=?',(u['id'],)).fetchone()[0]
        if not pwdcheck(x.old,old):raise HTTPException(400,'原密码不正确')
        c.execute('UPDATE users SET password=? WHERE id=?',(pwdhash(x.new),u['id']))
        c.execute('DELETE FROM sessions WHERE user_id=?',(u['id'],))
    return {'ok':True}

COURSES = [
    {'id':'writing','title':'学术论文写作','subtitle':'Academic Writing','description':'从研究问题、学术表达，到图表规范与论文修改。','state':'open','label':'课程已开放'},
    {'id':'english','title':'机械工程专业英语','subtitle':'English for Mechanical Engineering','description':'从感兴趣的话题开始，用简单英语对话，逐步建立口语学习画像。','state':'open','label':'英语对话已开放'},
    {'id':'equipment','title':'智能装备基础','subtitle':'Fundamentals of Intelligent Equipment','description':'面向智能装备的知识学习与实践，课程内容待建设。','state':'planned','label':'待建设'},
    {'id':'design','title':'智能制造课程设计','subtitle':'Intelligent Manufacturing Design','description':'通过AI五子棋交互案例，探索智能决策与工程实践。','state':'open','label':'实践案例已开放'},
]
@app.get('/api/courses')
def course_catalog(u=Depends(current)):
    return COURSES

GOMOKU_ROOT=ROOT/'course_assets'/'gomoku'
GOMOKU_FILES={'index.html','ai-worker.js','gomoku_uint8_qop_u8u8.onnx','gomoku_single.onnx','index_single_wasm118_mobile.html','index_single_wasm_abtest.html','gomoku.cpp','vct-worker.js','README.md'}
@app.get('/courses/design/gomoku/')
def gomoku_home(u=Depends(current)):
    path=GOMOKU_ROOT/'index.html'
    if not path.is_file():raise HTTPException(503,'课程案例资源未安装，请执行新版升级脚本')
    return FileResponse(path,media_type='text/html',headers={'Cache-Control':'no-store'})
@app.get('/courses/design/gomoku/{asset:path}')
def gomoku_asset(asset:str,u=Depends(current)):
    if asset not in GOMOKU_FILES:raise HTTPException(404,'案例资源不存在')
    path=GOMOKU_ROOT/asset
    if not path.is_file():raise HTTPException(404,'案例资源未安装')
    media='application/javascript' if asset.endswith('.js') else 'application/octet-stream' if asset.endswith('.onnx') else 'text/html' if asset.endswith('.html') else 'text/plain'
    return FileResponse(path,media_type=media,headers={'Cache-Control':'private, max-age=300'})

@app.get('/api/chapters')
def chapters(u=Depends(current)):
    with db() as c:
        rows=[dict(x) for x in c.execute('SELECT * FROM chapters')]
        for row in rows:row['materials']=[dict(x) for x in c.execute('SELECT id,title,filename,created FROM materials WHERE chapter_id=?',(row['id'],))]
    return rows

async def save_upload(file,extensions,limit):
    name=Path((file.filename or '').replace('\\','/')).name[:150]
    ext=Path(name).suffix.lower()
    if ext not in extensions:raise HTTPException(400,'不支持此文件格式')
    path=DATA/'files'/(secrets.token_hex(20)+ext)
    size=0
    try:
        with path.open('wb') as f:
            while chunk:=await file.read(65536):
                size+=len(chunk)
                if size>limit:raise HTTPException(413,'文件不能超过5 MB' if limit==LIMIT else f'文件不能超过 {limit//1024//1024} MiB')
                f.write(chunk)
        if not size:raise HTTPException(400,'不能上传空文件')
        if ext=='.pdf' and not path.read_bytes()[:5]==b'%PDF-':raise HTTPException(400,'不是有效PDF文件')
        if ext in ('.docx','.pptx'):
            try:
                with zipfile.ZipFile(path) as z:
                    if sum(i.file_size for i in z.infolist())>80*1024*1024:raise ValueError()
                    marker='word/document.xml' if ext=='.docx' else 'ppt/presentation.xml'
                    if marker not in z.namelist():raise ValueError()
            except Exception:raise HTTPException(400,'Office文件无效或解压后过大')
        return str(path),name
    except Exception:
        path.unlink(missing_ok=True);raise
    finally:await file.close()

@app.post('/api/materials')
async def material(chapter_id:int=Form(...),title:str=Form(...),file:UploadFile=File(...),u=Depends(teacher)):
    with db() as c:
        if not c.execute('SELECT 1 FROM chapters WHERE id=?',(chapter_id,)).fetchone():raise HTTPException(404,'章节不存在')
    if not 1<=len(title.strip())<=150:raise HTTPException(400,'课件名称应为1至150字')
    path,name=await save_upload(file,{'.pdf','.pptx','.ppt'},50*1024*1024)
    with db() as c:c.execute('INSERT INTO materials(chapter_id,title,path,filename,created) VALUES(?,?,?,?,?)',(chapter_id,title.strip(),path,name,now()))
    return {'ok':True}
@app.get('/api/materials/{mid}')
def download_material(mid:int,preview:bool=False,u=Depends(current)):
    with db() as c:m=c.execute('SELECT * FROM materials WHERE id=?',(mid,)).fetchone()
    if not m:raise HTTPException(404,'课件不存在')
    if preview and Path(m['filename']).suffix.lower()=='.pdf':return FileResponse(m['path'],media_type='application/pdf',headers={'Content-Disposition':'inline'})
    return FileResponse(m['path'],filename=m['filename'])
@app.delete('/api/materials/{mid}')
def delete_material(mid:int,u=Depends(teacher)):
    with db() as c:
        m=c.execute('SELECT * FROM materials WHERE id=?',(mid,)).fetchone()
        if not m:raise HTTPException(404,'课件不存在')
        c.execute('DELETE FROM materials WHERE id=?',(mid,))
    Path(m['path']).unlink(missing_ok=True)
    return {'ok':True}

def progress(uid,c):
    rows=c.execute('SELECT q.chapter_id,q.id,a.score FROM questions q LEFT JOIN attempts a ON a.id=(SELECT MIN(id) FROM attempts WHERE user_id=? AND question_id=q.id)',(uid,)).fetchall()
    result=[]
    for ch in c.execute('SELECT id,title FROM chapters'):
        items=[r for r in rows if r['chapter_id']==ch['id']]
        done=[r for r in items if r['score'] is not None]
        result.append({'id':ch['id'],'title':ch['title'],'total':len(items),'done':len(done),'score':round(sum(r['score'] for r in done)/len(items),1) if items else 0})
    return {'chapters':result,'done':sum(r['score'] is not None for r in rows),'total':len(rows),'score':round(sum(r['score'] or 0 for r in rows)/len(rows),1) if rows else 0}
@app.get('/api/progress')
def get_progress(u=Depends(current)):
    with db() as c:return progress(u['id'],c)
def learner_profile(uid,c,include_private=False):
    profile=c.execute('SELECT * FROM learner_profiles WHERE user_id=?',(uid,)).fetchone()
    result=dict(profile) if profile else {'user_id':uid,'research_area':'','writing_stage':'选题阶段','goal':'','teacher_note':'','updated':None}
    if not include_private:result.pop('teacher_note',None)
    attempts=[dict(r) for r in c.execute('SELECT a.*,q.chapter_id,q.level,q.stem,q.explanation FROM attempts a JOIN questions q ON q.id=a.question_id WHERE a.user_id=? ORDER BY a.id',(uid,))]
    modules=[]
    for ch in c.execute('SELECT id,title FROM chapters'):
        aa=[r for r in attempts if r['chapter_id']==ch['id']]
        first={}
        for r in aa:first.setdefault(r['question_id'],r)
        recent=aa[-8:]
        # Recent practice updates the learning estimate, but never rewrites assessed first-attempt scores.
        mastery=round(.35*sum(r['score'] for r in first.values())/len(first)+.65*sum(r['score'] for r in recent)/len(recent),1) if first else None
        target=max(1,min(3,aa[-1]['level']+(1 if aa[-1]['score']==100 else -1))) if aa else 1
        modules.append({'id':ch['id'],'title':ch['title'],'attempts':len(aa),'unique_questions':len(first),'first_accuracy':round(sum(r['score'] for r in first.values())/len(first),1) if first else None,'recent_accuracy':round(sum(r['score'] for r in recent)/len(recent),1) if recent else None,'recent_sample':len(recent),'mastery':mastery,'recommended_level':target,'last_practiced':aa[-1]['created'] if aa else None})
    weak=sorted([m for m in modules if m['mastery'] is not None and m['mastery']<70],key=lambda x:x['mastery'])
    unseen=[m for m in modules if m['mastery'] is None]
    rec=weak[:2] or unseen[:2] or sorted(modules,key=lambda x:x['mastery'])[:2]
    writing=[]
    for row in c.execute("SELECT s.id,s.ai_json,s.published,s.returned,s.created,s.status,a.title FROM submissions s JOIN assignments a ON a.id=s.assignment_id WHERE s.user_id=? ORDER BY s.created DESC LIMIT 5",(uid,)):
        if row['ai_json'] and (include_private or row['published'] or row['returned']):
            ai=json.loads(row['ai_json'])
            writing.append({'submission_id':row['id'],'title':row['title'],'created':row['created'],'source':'教师已发布的作业反馈' if row['published'] else ('已退回稿件的形式预审' if row['returned'] else 'AI初评，待教师复核'),'dimensions':ai.get('dimensions',[]),'triage':ai.get('triage',{}),'formal_check':ai.get('formal_check')})
    result.update(modules=modules,weak_topics=[m['title'] for m in weak],recommendations=[{'chapter_id':m['id'],'title':m['title'],'level':m['recommended_level'],'reason':('尚无作答记录，先完成基础诊断' if m['mastery'] is None else f"首次正确率 {m['first_accuracy']}%，近期{m['recent_sample']}次正确率 {m['recent_accuracy']}%，建议继续巩固") } for m in rec],recent_mistakes=[{'chapter_id':r['chapter_id'],'stem':r['stem'],'explanation':r['explanation'],'created':r['created']} for r in reversed(attempts) if r['score']==0][:8],writing_reviews=writing,total_attempts=len(attempts),assessment=progress(uid,c))
    return result
class ProfileUpdate(BaseModel):
    research_area:str=Field(default='',max_length=200)
    writing_stage:str=Field(default='选题阶段',max_length=50)
    goal:str=Field(default='',max_length=1000)
class TeacherProfileUpdate(ProfileUpdate):
    teacher_note:str=Field(default='',max_length=3000)
@app.get('/api/profile')
def get_profile(u=Depends(current)):
    with db() as c:return learner_profile(u['id'],c)
@app.post('/api/profile')
def update_profile(x:ProfileUpdate,u=Depends(current)):
    with db() as c:
        c.execute('INSERT INTO learner_profiles(user_id,research_area,writing_stage,goal,updated) VALUES(?,?,?,?,?) ON CONFLICT(user_id) DO UPDATE SET research_area=excluded.research_area,writing_stage=excluded.writing_stage,goal=excluded.goal,updated=excluded.updated',(u['id'],x.research_area,x.writing_stage,x.goal,now()))
    return {'ok':True}
@app.get('/api/teacher/students/{uid}/profile')
def teacher_profile(uid:int,u=Depends(teacher)):
    with db() as c:
        person=c.execute("SELECT id,name,student_no FROM users WHERE id=? AND role='student'",(uid,)).fetchone()
        if not person:raise HTTPException(404,'学生不存在')
        return {'student':dict(person),'profile':learner_profile(uid,c,True)}
@app.post('/api/teacher/students/{uid}/profile')
def teacher_update_profile(uid:int,x:TeacherProfileUpdate,u=Depends(teacher)):
    with db() as c:
        if not c.execute("SELECT 1 FROM users WHERE id=? AND role='student'",(uid,)).fetchone():raise HTTPException(404,'学生不存在')
        c.execute('INSERT INTO learner_profiles(user_id,research_area,writing_stage,goal,teacher_note,updated) VALUES(?,?,?,?,?,?) ON CONFLICT(user_id) DO UPDATE SET research_area=excluded.research_area,writing_stage=excluded.writing_stage,goal=excluded.goal,teacher_note=excluded.teacher_note,updated=excluded.updated',(uid,x.research_area,x.writing_stage,x.goal,x.teacher_note,now()))
    audit(u['id'],'update-profile',uid);return {'ok':True}

@app.get('/api/practice/next')
def next_question(chapter_id:int=0,u=Depends(current)):
    with db() as c:
        p=progress(u['id'],c)
        lp=learner_profile(u['id'],c)
        candidates=c.execute('SELECT * FROM questions WHERE (?=0 OR chapter_id=?)',(chapter_id,chapter_id)).fetchall()
        if not candidates:raise HTTPException(404,'章节不存在或没有题目')
        history={r['question_id']:dict(r) for r in c.execute('SELECT * FROM attempts WHERE user_id=? ORDER BY id',(u['id'],))}
        remaining=[r for r in candidates if r['id'] not in history]
        latest=c.execute('SELECT a.score,q.level FROM attempts a JOIN questions q ON q.id=a.question_id WHERE user_id=? ORDER BY a.id DESC LIMIT 1',(u['id'],)).fetchone()
        target=max(1,min(3,latest['level']+(1 if latest['score']==100 else -1))) if latest else 1
        if remaining:
            # Prioritize weaker chapters; difficulty adapts without changing common assessed pool.
            strengths={r['id']:(r['mastery'] if r['mastery'] is not None else 50) for r in lp['modules']}
            targets={r['id']:r['recommended_level'] for r in lp['modules']}
            q=min(remaining,key=lambda r:(strengths[r['chapter_id']],abs(r['level']-targets[r['chapter_id']]),r['id']))
            mode='首次作答计分'
        else:
            q=min(candidates,key=lambda r:(history[r['id']]['score'],history[r['id']]['id']))
            mode='巩固练习，不改变首次成绩'
        return {'id':q['id'],'chapter_id':q['chapter_id'],'level':q['level'],'stem':q['stem'],'options':json.loads(q['options']),'mode':mode,'progress':p,'recommendation_reason':next(r['reason'] for r in lp['recommendations'] if r['chapter_id']==q['chapter_id']) if any(r['chapter_id']==q['chapter_id'] for r in lp['recommendations']) else '结合个人档案与剩余题目推荐'}
class Answer(BaseModel):
    question_id:int
    choice:int=Field(ge=0,le=3)
@app.post('/api/practice/answer')
def answer(x:Answer,u=Depends(current)):
    throttle('answers:'+str(u['id']),120,60)
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        q=c.execute('SELECT * FROM questions WHERE id=?',(x.question_id,)).fetchone()
        if not q:raise HTTPException(404,'题目不存在')
        first=not c.execute('SELECT 1 FROM attempts WHERE user_id=? AND question_id=?',(u['id'],q['id'])).fetchone()
        score=100 if x.choice==q['answer'] else 0
        c.execute('INSERT INTO attempts(user_id,question_id,choice,score,created) VALUES(?,?,?,?,?)',(u['id'],q['id'],x.choice,score,now()))
    return {'score':score,'answer':q['answer'],'explanation':q['explanation'],'counted':first}

def submission_deletable(c,s):
    # A failed/interrupted AI attempt without results is still unreviewed.
    status=s['status'] or ''
    unreviewed=status in ('待批改','AI任务中断，可重新批改') or status.startswith('预审需人工处理：')
    return bool(unreviewed and not s['ai_json'] and s['score'] is None and not s['feedback']
                and not s['published'] and not s['returned'] and not s['return_reason']
                and not s['returned_at'] and not s['appeal_reason']
                and not c.execute('SELECT 1 FROM submission_history WHERE submission_id=? LIMIT 1',(s['id'],)).fetchone())

@app.get('/api/assignments')
def assignments(u=Depends(current)):
    with db() as c:
        result=[dict(r) for r in c.execute('SELECT * FROM assignments ORDER BY id DESC')]
        for a in result:
            a['rubric']=json.loads(a['rubric'])
            s=c.execute('SELECT id,filename,created,status,score,feedback,published,returned,return_reason,returned_at,appeal_reason,ai_json FROM submissions WHERE assignment_id=? AND user_id=?',(a['id'],u['id'])).fetchone()
            a['submission']=dict(s) if s else None
            if s:
                a['submission'].pop('ai_json')
                a['submission']['can_delete']=u['role']=='student' and submission_deletable(c,s)
            if s and not s['published']:
                a['submission'].update(score=None,feedback='')
                if not s['returned'] and s['status']=='AI建议退回，待教师确认':a['submission']['status']='AI已评，待教师复核'
    return result
class Criterion(BaseModel):
    name:str=Field(min_length=1,max_length=80)
    max:float=Field(gt=0,le=100)
class Assignment(BaseModel):
    title:str=Field(min_length=1,max_length=150)
    instructions:str=Field(default='',max_length=10000)
    deadline:str=Field(default='',max_length=50)
    grading_instructions:str=Field(default='',max_length=8000)
    rubric:list[Criterion]|None=Field(default=None,min_length=1,max_length=10)
    expected_revision:int|None=None
    workflow:bool=False

def assignment_values(x,c):
    if x.workflow and not x.instructions.strip():x.instructions='围绕“'+x.title.strip()+'”完成四轮递进写作。以本人真实材料或获准案例为依据，按每轮任务要求完善同一作品，不编造数据和文献。'
    if not x.title.strip() or not x.instructions.strip():raise HTTPException(400,'作业名称和要求不能为空')
    if x.deadline:
        try:
            if datetime.fromisoformat(x.deadline).tzinfo is not None:raise ValueError()
        except ValueError:raise HTTPException(400,'截止时间请使用北京时间，不带时区后缀')
    rubric=[r.model_dump() for r in x.rubric] if x.rubric is not None else None
    if rubric is not None:
        names=[r['name'].strip() for r in rubric]
        if len(set(names))!=len(names) or not all(names) or abs(sum(r['max'] for r in rubric)-100)>.001:raise HTTPException(400,'评分维度名称不能重复或为空，分值之和须为100')
        for item,name in zip(rubric,names):item['name']=name
    return rubric

@app.post('/api/assignments')
def create_assignment(x:Assignment,u=Depends(teacher)):
    with db() as c:
        rubric=assignment_values(x,c)
        if rubric is None:rubric=json.loads(c.execute('SELECT rubric FROM assignments ORDER BY id LIMIT 1').fetchone()[0])
        aid=c.execute('INSERT INTO assignments(title,instructions,rubric,deadline,created,grading_instructions) VALUES(?,?,?,?,?,?)',(x.title.strip(),x.instructions.strip(),json.dumps(rubric,ensure_ascii=False),x.deadline,now(),x.grading_instructions)).lastrowid
        c.execute('UPDATE assignments SET workflow=? WHERE id=?',(int(x.workflow),aid))
    audit(u['id'],'create-assignment',aid);return {'ok':True,'id':aid}

@app.put('/api/assignments/{aid}')
def edit_assignment(aid:int,x:Assignment,u=Depends(teacher)):
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        old=c.execute('SELECT * FROM assignments WHERE id=?',(aid,)).fetchone()
        if not old:raise HTTPException(404,'作业不存在')
        if bool(old['workflow'])!=x.workflow and c.execute('SELECT 1 FROM submissions WHERE assignment_id=?',(aid,)).fetchone():raise HTTPException(409,'已有提交的作业不能切换流程，请另建四轮作业')
        if x.expected_revision is not None and x.expected_revision!=old['revision']:raise HTTPException(409,'作业已被更新，请刷新后再编辑')
        if c.execute("SELECT 1 FROM submissions WHERE assignment_id=? AND status IN ('AI排队中','AI批改中') LIMIT 1",(aid,)).fetchone():raise HTTPException(409,'本作业正在AI批改，请完成或取消排队任务后再编辑')
        rubric=assignment_values(x,c)
        c.execute('UPDATE assignments SET title=?,instructions=?,deadline=?,grading_instructions=?,rubric=?,revision=revision+1 WHERE id=?',(x.title.strip(),x.instructions.strip(),x.deadline,x.grading_instructions,json.dumps(rubric,ensure_ascii=False) if rubric is not None else old['rubric'],aid))
        c.execute('UPDATE assignments SET workflow=? WHERE id=?',(int(x.workflow),aid))
    audit(u['id'],'edit-assignment',aid);return {'ok':True,'revision':old['revision']+1}
@app.post('/api/submissions')
async def submit(assignment_id:int=Form(...),file:UploadFile=File(...),u=Depends(current)):
    throttle('upload:'+str(u['id']),20,3600)
    if u['role']!='student':raise HTTPException(400,'请使用学生账号提交作业')
    with db() as c:
        a=c.execute('SELECT * FROM assignments WHERE id=?',(assignment_id,)).fetchone()
        if not a:raise HTTPException(404,'作业不存在')
        if a['workflow']:raise HTTPException(409,'请使用四轮写作任务页面提交')
        old=c.execute('SELECT * FROM submissions WHERE assignment_id=? AND user_id=?',(assignment_id,u['id'])).fetchone()
        if old and not old['returned']:raise HTTPException(409,'此作业已提交；如需重交请联系教师退回')
        # Returned submissions may be revised after the original deadline.
        if not (old and old['returned']) and a['deadline'] and datetime.now()>datetime.fromisoformat(a['deadline']):raise HTTPException(400,'已超过作业截止时间')
    path,name=await save_upload(file,{'.docx','.pdf'},LIMIT)
    cfg=ai_settings();auto=bool(cfg.get('key') and cfg.get('model') and cfg.get('auto_review',True))
    try:
        with db() as c:
            c.execute('BEGIN IMMEDIATE')
            old=c.execute('SELECT * FROM submissions WHERE assignment_id=? AND user_id=?',(assignment_id,u['id'])).fetchone()
            status='待批改'
            if old:
                if not old['returned']:raise HTTPException(409,'作业已提交')
                c.execute('INSERT INTO submission_history(submission_id,user_id,path,filename,snapshot,archived) VALUES(?,?,?,?,?,?)',(old['id'],u['id'],old['path'],old['filename'],json.dumps(dict(old),ensure_ascii=False),now()))
                sid=old['id']
                c.execute("UPDATE submissions SET path=?,filename=?,created=?,status=?,ai_json=NULL,score=NULL,feedback='',published=0,returned=0,return_reason='',returned_at='',appeal_reason='' WHERE id=?",(path,name,now(),status,sid))
            else:
                # Never reuse a removed submission ID: stale teacher pages and old
                # usage records must not point at a different student's new upload.
                sid=c.execute("SELECT MAX(value)+1 FROM (SELECT COALESCE(MAX(id),0) value FROM submissions UNION ALL SELECT COALESCE(MAX(CAST(target AS INTEGER)),0) value FROM audit WHERE action='student-delete-submission')").fetchone()[0]
                c.execute('INSERT INTO submissions(id,assignment_id,user_id,path,filename,created,status) VALUES(?,?,?,?,?,?,?)',(sid,assignment_id,u['id'],path,name,now(),status))
    except Exception:
        Path(path).unlink(missing_ok=True);raise
    if auto:
        due=next_daily(cfg.get('review_time','00:30')) if cfg.get('review_mode')=='scheduled' else None
        try:REVIEW_QUEUE.enqueue(sid,run_at=due)
        except HTTPException:auto=False
    return {'ok':True,'auto_review':auto}
@app.delete('/api/submissions/{sid}')
def delete_submission(sid:int,u=Depends(current)):
    if u['role']!='student':raise HTTPException(403,'请使用学生账号删除自己的未批改作业')
    with db() as c:
        # Serializes against teacher grading, return and AI queue admission.
        c.execute('BEGIN IMMEDIATE')
        s=c.execute('SELECT * FROM submissions WHERE id=? AND user_id=?',(sid,u['id'])).fetchone()
        if not s:raise HTTPException(404,'作业不存在')
        if c.execute('SELECT 1 FROM writing_rounds WHERE submission_id=?',(sid,)).fetchone():raise HTTPException(409,'四轮写作请使用任务页替换未检查的文件，历史版本需要保留')
        if not submission_deletable(c,s):raise HTTPException(409,'作业已批改、正在批改或存在评阅历史，不能自行删除；如需修改请联系教师退回')
        c.execute('DELETE FROM submissions WHERE id=?',(sid,))
        c.execute('INSERT INTO audit(user_id,action,target,created) VALUES(?,?,?,?)',(u['id'],'student-delete-submission',str(sid),now()))
    # Uploads have random server-managed names; do not remove unrelated paths.
    path=Path(s['path']).resolve()
    if path.is_relative_to((DATA/'files').resolve()):
        try:path.unlink(missing_ok=True)
        except OSError:
            import logging
            logging.getLogger(__name__).warning('Deleted submission %s has an upload pending cleanup',sid)
    return {'ok':True}
@app.get('/api/submissions/{sid}/file')
def submission_file(sid:int,u=Depends(current)):
    with db() as c:s=c.execute('SELECT * FROM submissions WHERE id=?',(sid,)).fetchone()
    if not s:raise HTTPException(404,'记录不存在')
    if u['role']!='teacher' and s['user_id']!=u['id']:raise HTTPException(403,'无权读取其他同学的作业')
    return FileResponse(s['path'],filename=s['filename'])
@app.get('/api/teacher/submissions')
def all_submissions(u=Depends(teacher)):
    with db() as c:
        rows=[dict(r) for r in c.execute('SELECT s.id,s.assignment_id,s.user_id,s.filename,s.created,s.status,s.ai_json,s.score,s.feedback,s.published,s.returned,s.return_reason,s.returned_at,s.appeal_reason,u.name,u.student_no,a.title,a.workflow,a.revision assignment_revision FROM submissions s JOIN users u ON u.id=s.user_id JOIN assignments a ON a.id=s.assignment_id ORDER BY s.id DESC')]
        for r in rows:r['ai']=json.loads(r.pop('ai_json')) if r['ai_json'] else None
    return rows
class Grade(BaseModel):
    score:float=Field(ge=0,le=100)
    feedback:str=Field(max_length=20000)
    published:bool=False
@app.post('/api/teacher/submissions/{sid}/grade')
def grade(sid:int,x:Grade,u=Depends(teacher)):
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        if not c.execute('SELECT 1 FROM submissions WHERE id=?',(sid,)).fetchone():raise HTTPException(404,'作业不存在')
        writing_workflow.require_final(c,sid)
        if c.execute('SELECT status FROM submissions WHERE id=?',(sid,)).fetchone()[0] in BUSY:raise HTTPException(409,'请等待AI任务结束后再保存评分')
        if c.execute('SELECT returned FROM submissions WHERE id=?',(sid,)).fetchone()[0]:raise HTTPException(409,'已退回的作业须等待学生重交后评分')
        c.execute('UPDATE submissions SET score=?,feedback=?,published=?,status=? WHERE id=?',(x.score,x.feedback,int(x.published),'已发布' if x.published else '教师复核中',sid))
    audit(u['id'],'grade',sid)
    return {'ok':True}
class ReturnWork(BaseModel):
    reason:str=Field(min_length=5,max_length=5000)
@app.post('/api/teacher/submissions/{sid}/return')
def return_work(sid:int,x:ReturnWork,u=Depends(teacher)):
    with db() as c:
        s=c.execute('SELECT * FROM submissions WHERE id=?',(sid,)).fetchone()
        if not s:raise HTTPException(404,'作业不存在')
        if s['status'] in BUSY:raise HTTPException(409,'请等待AI任务结束后再退回')
        if s['returned']:raise HTTPException(409,'该作业已退回，等待学生修改')
        c.execute("UPDATE submissions SET returned=1,return_reason=?,returned_at=?,published=0,status='已退回，待修改' WHERE id=?",(x.reason.strip(),now(),sid))
    audit(u['id'],'return-for-revision',sid);return {'ok':True}
@app.delete('/api/teacher/submissions/{sid}')
def return_submission(sid:int,u=Depends(teacher)):
    raise HTTPException(400,'请使用退回接口并填写原因；作业原文与记录将保留')
@app.post('/api/teacher/submissions/{sid}/reopen')
def reopen(sid:int,u=Depends(teacher)):
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        row=c.execute('SELECT * FROM submissions WHERE id=?',(sid,)).fetchone()
        if not row:raise HTTPException(404,'作业不存在')
        if not row['returned']:raise HTTPException(409,'当前作业未退回')
        c.execute("UPDATE submissions SET returned=0,status='教师恢复人工复核' WHERE id=?",(sid,))
    audit(u['id'],'reopen-returned',sid);return {'ok':True}
@app.get('/api/teacher/submissions/{sid}/history')
def submission_history(sid:int,u=Depends(teacher)):
    with db() as c:
        rows=[]
        for r in c.execute('SELECT * FROM submission_history WHERE submission_id=? ORDER BY id DESC',(sid,)):
            snap=json.loads(r['snapshot']);rows.append({'id':r['id'],'filename':r['filename'],'created':snap['created'],'archived':r['archived'],'return_reason':snap.get('return_reason',''),'score':snap.get('score'),'ai':json.loads(snap['ai_json']) if snap.get('ai_json') else None})
    return rows
@app.get('/api/teacher/history/{hid}/file')
def history_file(hid:int,u=Depends(teacher)):
    with db() as c:r=c.execute('SELECT * FROM submission_history WHERE id=?',(hid,)).fetchone()
    if not r:raise HTTPException(404,'版本不存在')
    return FileResponse(r['path'],filename=r['filename'])
@app.get('/api/teacher/students')
def students(u=Depends(teacher)):
    with db() as c:
        result=[dict(r) for r in c.execute("SELECT id,name,student_no,created FROM users WHERE role='student' ORDER BY student_no")]
        for r in result:
            r['practice']=progress(r['id'],c)
            r['submitted']=c.execute('SELECT COUNT(*) FROM submissions WHERE user_id=?',(r['id'],)).fetchone()[0]
            r['assignment_score']=c.execute('SELECT AVG(score) FROM submissions WHERE user_id=? AND published=1',(r['id'],)).fetchone()[0]
    return result
class Reset(BaseModel):password:str=Field(min_length=10,max_length=128)
@app.post('/api/teacher/students/{uid}/password')
def reset(uid:int,x:Reset,u=Depends(teacher)):
    with db() as c:
        if not c.execute("SELECT 1 FROM users WHERE id=? AND role='student'",(uid,)).fetchone():raise HTTPException(404,'学生不存在')
        c.execute('UPDATE users SET password=? WHERE id=?',(pwdhash(x.password),uid))
        c.execute('DELETE FROM sessions WHERE user_id=?',(uid,))
    audit(u['id'],'reset-password',uid);return {'ok':True}
@app.get('/api/teacher/export')
def export(u=Depends(teacher)):
    rows=students(u)
    f=io.StringIO();f.write('\ufeff');w=csv.writer(f)
    w.writerow(['学号','姓名','练习完成题数','统一题库题数','练习平时成绩(未答计0)','已交作业数','已发布作业平均分'])
    for r in rows:w.writerow([("'"+r['student_no']),r['name'] if not r['name'].startswith(('=','+','-','@')) else "'"+r['name'],r['practice']['done'],r['practice']['total'],r['practice']['score'],r['submitted'],r['assignment_score']])
    return StreamingResponse(iter([f.getvalue()]),media_type='text/csv; charset=utf-8',headers={'Content-Disposition':'attachment; filename=grades.csv'})

def ai_settings():
    path=DATA/'ai.json'
    return json.loads(path.read_text()) if path.exists() else {'base_url':'https://api.openai.com/v1','model':'','key':''}
@app.get('/api/teacher/ai-settings')
def get_ai(u=Depends(teacher)):
    s=ai_settings();return {'base_url':s['base_url'],'model':s['model'],'configured':bool(s.get('key') and s.get('model')),'auto_review':s.get('auto_review',True),'auto_return':s.get('auto_return',True),'error_threshold':s.get('error_threshold',5),'formal_standard':s.get('formal_standard',DEFAULT_STANDARD),'review_mode':s.get('review_mode','immediate' if s.get('auto_review',True) else 'manual'),'review_time':s.get('review_time','00:30'),'off_peak_only':s.get('off_peak_only',False),'review_version':'1.5.1'}
class AISettings(BaseModel):
    review_mode:str|None=Field(default=None,pattern='^(immediate|manual|scheduled)$')
    review_time:str=Field(default='00:30',pattern=r'^([01]\d|2[0-3]):[0-5]\d$')
    off_peak_only:bool=False
    auto_review:bool=True
    auto_return:bool=True
    error_threshold:int=Field(default=5,ge=3,le=20)
    formal_standard:str=Field(default=DEFAULT_STANDARD,min_length=20,max_length=8000)
    base_url:str=Field(max_length=500)
    model:str=Field(max_length=150)
    key:str=Field(default='',max_length=500)
@app.post('/api/teacher/ai-settings')
def set_ai(x:AISettings,u=Depends(teacher)):
    parsed=urlparse(x.base_url)
    if parsed.scheme!='https' or not parsed.hostname or parsed.username or parsed.password or parsed.query or parsed.fragment:raise HTTPException(400,'接口地址须为有效HTTPS基础地址，例如 https://api.openai.com/v1')
    old=ai_settings();mode=x.review_mode or ('immediate' if x.auto_review else 'manual');data={'base_url':x.base_url.rstrip('/'),'model':x.model.strip(),'key':x.key.strip() or old.get('key',''),'auto_review':mode!='manual','review_mode':mode,'review_time':x.review_time,'off_peak_only':x.off_peak_only,'auto_return':x.auto_return,'error_threshold':x.error_threshold,'formal_standard':x.formal_standard}
    tmp=DATA/'ai.tmp';fd=os.open(tmp,os.O_WRONLY|os.O_CREAT|os.O_TRUNC,0o600)
    with os.fdopen(fd,'w') as f:json.dump(data,f)
    tmp.replace(DATA/'ai.json');audit(u['id'],'ai-settings','updated')
    return {'ok':True,'configured':bool(data['key'] and data['model'])}

def ai_call(system,payload):
    import base64
    cfg=ai_settings()
    if not cfg.get('key') or not cfg.get('model'):raise ValueError('尚未配置AI接口')
    body=dict(payload);images=body.pop('_images',[]);ctx=body.pop('_usage_context',{})
    text=json.dumps(body,ensure_ascii=False)
    # Cache is scoped to a student and the exact prompt, standard, model and page bytes.
    digest=hashlib.sha256((cfg['base_url']+cfg['model']+str(ctx.get('user_id',''))+system+text).encode())
    for path in images:digest.update(Path(path).read_bytes())
    key=digest.hexdigest()
    with db() as c:cached=c.execute('SELECT payload FROM ai_result_cache WHERE hash=?',(key,)).fetchone()
    def log(usage,local):
        with db() as c:c.execute('INSERT INTO ai_usage(user_id,submission_id,purpose,model,input_tokens,output_tokens,cache_hit_tokens,local_cache,created) VALUES(?,?,?,?,?,?,?,?,?)',(ctx.get('user_id'),ctx.get('submission_id'),ctx.get('purpose','练习'),cfg['model'],int(usage.get('prompt_tokens',0)),int(usage.get('completion_tokens',0)),int(usage.get('prompt_cache_hit_tokens',0)),int(local),now()))
    if cached and not ctx.get('refresh',False):
        log({},True);return json.loads(cached['payload'])
    content=[{'type':'text','text':text}]
    for path in images:
        content.append({'type':'image_url','image_url':{'url':'data:image/jpeg;base64,'+base64.b64encode(Path(path).read_bytes()).decode(),'detail':'original'}})
    request={'model':cfg['model'],'messages':[{'role':'system','content':system},{'role':'user','content':content if images else text}],'response_format':{'type':'json_object'},'max_tokens':5000 if images else 2000}
    if urlparse(cfg['base_url']).hostname=='api.deepseek.com':request['thinking']={'type':'disabled'}
    with httpx.Client(timeout=180,follow_redirects=False) as client:
        r=client.post(cfg['base_url']+'/chat/completions',headers={'Authorization':'Bearer '+cfg['key']},json=request)
        if r.status_code>=400:raise ValueError(f'AI服务返回 HTTP {r.status_code}，请检查接口、视觉模型或余额')
        data=r.json();log(data.get('usage',{}),False)
        try:
            if data['choices'][0].get('finish_reason')=='length':raise ValueError('AI输出超限，请教师复核或调整任务范围')
            result=json.loads(data['choices'][0]['message']['content'])
            if not isinstance(result,dict):raise ValueError()
        except Exception:raise ValueError('AI返回格式不完整，请教师复核；未自动退回')
    with db() as c:c.execute('INSERT OR REPLACE INTO ai_result_cache VALUES(?,?,?)',(key,json.dumps(result,ensure_ascii=False),now()))
    return result

@app.get('/api/teacher/ai-usage')
def ai_usage(u=Depends(teacher)):
    with db() as c:
        rows=[dict(r) for r in c.execute('SELECT * FROM ai_usage ORDER BY id DESC LIMIT 100')]
        totals=dict(c.execute('SELECT COALESCE(SUM(input_tokens),0) input_tokens,COALESCE(SUM(output_tokens),0) output_tokens,COALESCE(SUM(local_cache),0) local_reuses,COUNT(*) requests FROM ai_usage').fetchone())
    return {'totals':totals,'records':rows}

def extract(path):
    p=Path(path)
    if p.suffix=='.docx':
        from docx import Document
        d=Document(p);text='\n'.join([r.text for r in d.paragraphs]+[' | '.join(c.text for c in r.cells) for t in d.tables for r in t.rows])
    else:
        from pypdf import PdfReader
        r=PdfReader(p)
        if len(r.pages)>150:raise ValueError('PDF超过150页，请拆分作业')
        text='\n'.join(page.extract_text() or '' for page in r.pages)
    if len(text.strip())<50:raise ValueError('文档可提取文字不足；扫描PDF请转换为文字型PDF或DOCX')
    if len(text)>60000:raise ValueError('文档超过60000字符，请拆分后提交；未进行截断评分')
    return text

def run_review(sid):
    with db() as c:
        workflow=c.execute('SELECT a.workflow FROM assignments a JOIN submissions s ON s.assignment_id=a.id WHERE s.id=?',(sid,)).fetchone()
    if workflow and workflow['workflow']:return writing_workflow.run_review(sid)
    try:
        with db() as c:
            s=dict(c.execute('SELECT s.*,a.instructions,a.rubric,a.title assignment_title,a.grading_instructions,a.revision assignment_revision FROM submissions s JOIN assignments a ON a.id=s.assignment_id WHERE s.id=?',(sid,)).fetchone())
            c.execute("UPDATE submissions SET status='AI批改中' WHERE id=?",(sid,))
        cfg=ai_settings()
        previous=json.loads(s['ai_json']) if s['ai_json'] else {}
        refresh_score=bool(s['ai_json'] and previous.get('score') is None)
        if urlparse(cfg['base_url']).hostname=='api.deepseek.com' and cfg.get('model')!='deepseek-flash':
            raise ValueError('图文预审请使用支持视觉的 deepseek-flash 模型')
        with prepare_document(s['path']) as document:
            # Rendering may cross the pricing boundary; check again before API use.
            if cfg.get('off_peak_only',False) and peak(time.time()):raise DeferReview()
            response=ai_call(FORMAL_PROMPT+'\n'+SCORING_PROMPT,{'grading_instructions':s['grading_instructions'],'assignment':s['instructions'],'rubric':json.loads(s['rubric']),'standard':cfg.get('formal_standard',DEFAULT_STANDARD),'document':document['text'],'page_count':document['page_count'],'warnings':document['warnings'],'_images':[p['image_path'] for p in document['pages']],'_usage_context':{'user_id':s['user_id'],'submission_id':sid,'purpose':'论文图文预审与建议评分','refresh':refresh_score}})
            formal=normalize_formal(response.get('formal_check'),document,cfg.get('error_threshold',5))
            scoring=normalize_score(response.get('dimensions'),json.loads(s['rubric']),document,formal)
        triage={'recommendation':'return' if formal['decision']=='revise' else 'review','reasons':[x['problem'] for x in formal['issues']],'required_changes':[x['change'] for x in formal['issues']],'evidence':[x['quote'] or x['evidence'] for x in formal['issues']],'evidence_matches':[x['evidence_located'] for x in formal['issues']]}
        result={'formal_check':formal,'triage':triage,'dimensions':scoring['dimensions'],'score':scoring['score'],'score_warning':scoring['warning'],'scoring_version':'1.5.1','scoring_basis':{'revision':s['assignment_revision'],'title':s['assignment_title'],'instructions':s['instructions'],'grading_instructions':s['grading_instructions'],'rubric':json.loads(s['rubric'])},'summary':str(response.get('summary',''))[:6000],'issues':[],'reference_note':str(response.get('reference_note','仅检查引用格式，未核实文献真实性'))[:2000],'reviewed_at':now()}
        auto=bool(cfg.get('auto_return',True) and formal['automatic_eligible'])
        if auto:status='AI形式预审退回，待修改'
        elif formal['decision']=='pass':status='形式预审通过，待教师评阅'
        else:status='形式预审需教师复核'
        with db() as c:
            c.execute('BEGIN IMMEDIATE')
            current=c.execute('SELECT path,published,returned FROM submissions WHERE id=?',(sid,)).fetchone()
            if not current or current['path']!=s['path'] or current['published'] or current['returned']:return
            c.execute('UPDATE submissions SET ai_json=?,status=? WHERE id=?',(json.dumps(result,ensure_ascii=False),status,sid))
            if auto:c.execute('UPDATE submissions SET returned=1,return_reason=?,returned_at=?,published=0 WHERE id=?',(return_reason(formal),now(),sid))
        audit(s['user_id'],'auto-formal-return' if auto else 'formal-check',sid)
        return True
    except DeferReview:raise
    except Exception as e:
        message=str(e) if isinstance(e,ValueError) else '文档解析或网络请求失败，请教师复核'
        with db() as c:c.execute('UPDATE submissions SET status=? WHERE id=?',('预审需人工处理：'+message[:180],sid))
        return False

class Appeal(BaseModel):reason:str=Field(min_length=5,max_length=2000)
@app.post('/api/submissions/{sid}/appeal')
def appeal(sid:int,x:Appeal,u=Depends(current)):
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        row=c.execute('SELECT * FROM submissions WHERE id=?',(sid,)).fetchone()
        if not row or row['user_id']!=u['id']:raise HTTPException(404,'作业不存在')
        if not row['returned']:raise HTTPException(409,'当前作业不在退回状态')
        c.execute("UPDATE submissions SET returned=0,appeal_reason=?,status='学生申请教师复核' WHERE id=?",(x.reason,sid))
    audit(u['id'],'appeal-formal-check',sid);return {'ok':True}

@app.post('/api/teacher/submissions/{sid}/ai')
def review(sid:int,u=Depends(teacher)):
    cfg=ai_settings()
    if not cfg.get('key') or not cfg.get('model'):raise HTTPException(400,'请先配置AI接口和模型')
    throttle('review:'+str(u['id']),40,3600)
    REVIEW_QUEUE.enqueue(sid,replace_ai=True);audit(u['id'],'ai-review',sid)
    return {'ok':True}

class BatchReview(BaseModel):
    assignment_id:int|None=None
    run_at:str=Field(default='',max_length=50)

@app.post('/api/teacher/reviews/batch')
def batch_review(x:BatchReview,u=Depends(teacher)):
    cfg=ai_settings()
    if not cfg.get('key') or not cfg.get('model'):raise HTTPException(400,'请先配置AI接口')
    throttle('batch-review:'+str(u['id']),12,3600)
    due=None
    if x.run_at:
        try:
            dt=datetime.fromisoformat(x.run_at)
            if dt.tzinfo is not None:raise ValueError()
            due=dt.replace(tzinfo=BJT).timestamp()
            if not time.time()+5<=due<=time.time()+366*86400:raise ValueError()
        except ValueError:raise HTTPException(400,'请选择未来一年内的北京时间，至少晚于当前5秒')
    with db() as c:
        rows=c.execute('SELECT id FROM submissions WHERE (? IS NULL OR assignment_id=?) ORDER BY id',(x.assignment_id,x.assignment_id)).fetchall()
    queued=0
    for row in rows:
        try:REVIEW_QUEUE.enqueue(row['id'],run_at=due,wake=False);queued+=1
        except HTTPException as e:
            if e.status_code not in (404,409):raise
    if due is None:REVIEW_QUEUE.wake()
    audit(u['id'],'batch-ai-review',f'{queued} queued')
    return {'ok':True,'queued':queued,'skipped':len(rows)-queued,'run_at':due,'off_peak_only':cfg.get('off_peak_only',False)}

@app.get('/api/teacher/review-jobs')
def review_jobs(u=Depends(teacher)):
    with db() as c:
        counts={r['state']:r['n'] for r in c.execute('SELECT state,COUNT(*) n FROM paper_review_jobs GROUP BY state')}
        rows=[dict(r) for r in c.execute('SELECT j.id,j.submission_id,j.run_at,j.state,j.error,j.created,u.name,u.student_no,a.title FROM paper_review_jobs j LEFT JOIN submissions s ON s.id=j.submission_id LEFT JOIN users u ON u.id=s.user_id LEFT JOIN assignments a ON a.id=s.assignment_id ORDER BY j.id DESC LIMIT 100')]
    return {'counts':counts,'jobs':rows,'timezone':'Asia/Shanghai','review_version':'1.5.1'}

@app.delete('/api/teacher/review-jobs/{jid}')
def cancel_review(jid:int,u=Depends(teacher)):
    REVIEW_QUEUE.cancel(jid);audit(u['id'],'cancel-ai-review',jid);return {'ok':True}

@app.post('/api/practice/ai')
def ai_practice(chapter_id:int=0,u=Depends(current)):
    throttle('ai-practice:'+str(u['id']),5,86400)
    with db() as c:
        p=progress(u['id'],c)
        lp=learner_profile(u['id'],c)
        chosen=chapter_id or lp['recommendations'][0]['chapter_id']
        ch=next((ch for ch in lp['modules'] if ch['id']==chosen),None)
        if not ch:raise HTTPException(404,'学习单元不存在')
        mistakes=[dict(x) for x in c.execute('SELECT q.stem,q.explanation FROM attempts a JOIN questions q ON q.id=a.question_id WHERE a.user_id=? AND a.score=0 AND q.chapter_id=? ORDER BY a.id DESC LIMIT 3',(u['id'],ch['id']))]
    result=ai_call('你是论文写作教师。依据薄弱知识点生成一道中文单选变式题，四个选项且仅一个正确，避免复述原题。返回JSON：stem字符串，options四字符串数组，answer为0至3整数，explanation字符串。学生仅通过作答后查看解析。',{'chapter':ch['title'],'mastery':ch['mastery'],'recommended_level':ch['recommended_level'],'research_area':lp['research_area'],'writing_stage':lp['writing_stage'],'learning_goal':lp['goal'],'mistakes':mistakes,'_usage_context':{'user_id':u['id'],'purpose':'个性化练习'}})
    if not isinstance(result.get('options'),list) or len(result['options'])!=4 or not all(isinstance(x,str) for x in result['options']) or type(result.get('answer')) is not int or result['answer'] not in range(4) or not isinstance(result.get('stem'),str) or not isinstance(result.get('explanation'),str):raise HTTPException(502,'AI题目格式不正确，请重试')
    with db() as c:pid=c.execute('INSERT INTO ai_practice(user_id,chapter_id,payload,created) VALUES(?,?,?,?)',(u['id'],ch['id'],json.dumps(result,ensure_ascii=False),now())).lastrowid
    return {'id':pid,'chapter_id':ch['id'],'stem':result['stem'],'options':result['options'],'mode':'AI个性化拓展，不计统一成绩'}
class AIAnswer(BaseModel):choice:int=Field(ge=0,le=3)
@app.post('/api/practice/ai/{pid}/answer')
def ai_answer(pid:int,x:AIAnswer,u=Depends(current)):
    with db() as c:
        r=c.execute('SELECT * FROM ai_practice WHERE id=? AND user_id=?',(pid,u['id'])).fetchone()
        if not r:raise HTTPException(404,'题目不存在')
        if r['choice'] is not None:raise HTTPException(409,'已经作答')
        q=json.loads(r['payload']);c.execute('UPDATE ai_practice SET choice=? WHERE id=?',(x.choice,pid))
    return {'score':100 if x.choice==q['answer'] else 0,'answer':q['answer'],'explanation':q['explanation'],'counted':False}
@app.exception_handler(ValueError)
async def value_error(request,exc):return JSONResponse({'detail':str(exc)},400)
@app.get('/healthz')
def health():return {'ok':True}
REVIEW_QUEUE=ReviewQueue(db,POOL,lambda sid:run_review(sid),ai_settings)
writing_workflow.register(app,globals())

import english_course
english_course.register(app,db,current,teacher,lambda *args: ai_call(*args),ai_settings,throttle,DATA)

app.mount('/static',StaticFiles(directory=ROOT/'static'),name='static')
@app.get('/')
def index():return FileResponse(ROOT/'static'/'index.html')
