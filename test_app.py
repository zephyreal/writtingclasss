import os,tempfile,io,json,time
os.environ['ENGLISH_NEWS_AUTO']='0'
import pytest
os.environ['DATA_DIR']=tempfile.mkdtemp(prefix='writing-test-')
from fastapi.testclient import TestClient
import app as a
from docx import Document

H={'X-App-Request':'1'}
def client():return TestClient(a.app,headers=H)
def invitation_for_tests():
    code='test-invite-only'
    with a.db() as db:db.execute('INSERT OR IGNORE INTO course_invites(code_hash,label,max_uses,expires,created) VALUES(?,?,?,?,?)',(a.hashlib.sha256(code.encode()).hexdigest(),'测试专用',10000,time.time()+86400,a.now()))
    return code
def student(no):
    c=client();assert c.post('/api/register',json={'name':'测试学生','student_no':no,'password':'test-password-123','invite_code':invitation_for_tests()}).status_code==200
    assert c.post('/api/login',json={'student_no':no,'password':'test-password-123','invite_code':invitation_for_tests()}).status_code==200
    return c
def teacher():
    with a.db() as db:
        if not db.execute("SELECT 1 FROM users WHERE student_no='teacher'").fetchone():db.execute("INSERT INTO users(name,student_no,password,role,created) VALUES(?,?,?,'teacher',?)",('教师','teacher',a.pwdhash('teacher-password-123'),a.now()))
    c=client();assert c.post('/api/login',json={'student_no':'teacher','password':'teacher-password-123'}).status_code==200
    return c
def doc():
    d=Document();d.add_paragraph('本文研究激光焊接速度与熔深的关系。采用三组实验测试不同速度对熔深的影响，以截面金相结果作为评价依据。结果表明在本实验条件下熔深存在差异，适用范围有待更多试验检验。')
    f=io.BytesIO();d.save(f);return f.getvalue()
def submit(c):return c.post('/api/submissions',data={'assignment_id':1},files={'file':('draft.docx',doc())})

def test_auth_csrf_roles():
    c=TestClient(a.app)
    assert c.get('/api/chapters').status_code==401
    assert c.post('/api/login',json={'student_no':'x','password':'x'}).status_code==403
    s=student('auth100');assert s.get('/api/teacher/students').status_code==403
    assert s.post('/api/logout',headers={'Origin':'https://evil.example'}).status_code==403
    cookie=s.cookies.get('session');assert cookie
    assert 'httponly' in s.post('/api/login',json={'student_no':'auth100','password':'test-password-123'}).headers['set-cookie'].lower()

def test_registration_validation_and_duplicate():
    s=student('reg100')
    assert s.post('/api/register',json={'name':'学生','student_no':'reg100','password':'test-password-123','invite_code':invitation_for_tests()}).status_code==409
    assert s.post('/api/register',json={'name':'学生','student_no':'bad;id','password':'test-password-123'}).status_code==400
    assert s.post('/api/register',json={'name':'学生','student_no':'reg200','password':'short'}).status_code==422

def test_student_delete_unreviewed_submission_and_resubmit(monkeypatch):
    from pathlib import Path
    monkeypatch.setattr(a,'ai_settings',lambda:{'key':'','model':''})
    s=student('delete100');other=student('delete200');t=teacher()
    assert submit(s).status_code==200
    sub=s.get('/api/assignments').json()[0]['submission'];sid=sub['id']
    assert sub['can_delete'] and 'ai_json' not in sub
    with a.db() as c:path=c.execute('SELECT path FROM submissions WHERE id=?',(sid,)).fetchone()[0]
    assert Path(path).exists()
    assert other.delete(f'/api/submissions/{sid}').status_code==404
    assert t.delete(f'/api/submissions/{sid}').status_code==403
    assert s.delete(f'/api/submissions/{sid}').status_code==200
    assert not Path(path).exists()
    assert s.get('/api/assignments').json()[0]['submission'] is None
    assert s.get(f'/api/submissions/{sid}/file').status_code==404
    assert s.delete(f'/api/submissions/{sid}').status_code==404
    assert submit(s).status_code==200
    assert s.get('/api/assignments').json()[0]['submission']['id']>sid
    assert t.post(f'/api/teacher/submissions/{sid}/grade',json={'score':90,'feedback':'旧页面','published':True}).status_code==404
    with a.db() as c:assert c.execute("SELECT 1 FROM audit WHERE action='student-delete-submission' AND target=?",(str(sid),)).fetchone()

def test_student_delete_blocks_review_states_and_history(monkeypatch):
    monkeypatch.setattr(a,'ai_settings',lambda:{'key':'','model':''})
    s=student('deleteprotected100');assert submit(s).status_code==200
    sid=s.get('/api/assignments').json()[0]['submission']['id']
    states=[{'status':'AI排队中'},{'status':'AI批改中'}, {'ai_json':'{}'},
            {'score':0},{'feedback':'教师草稿'}, {'published':1}, {'returned':1},
            {'return_reason':'修改'}, {'returned_at':'2026-10-19'},
            {'appeal_reason':'申请复核'}, {'status':'教师恢复人工复核'}]
    for changes in states:
        with a.db() as c:
            c.execute("UPDATE submissions SET status='待批改',ai_json=NULL,score=NULL,feedback='',published=0,returned=0,return_reason='',returned_at='',appeal_reason='' WHERE id=?",(sid,))
            c.execute('UPDATE submissions SET '+','.join(k+'=?' for k in changes)+' WHERE id=?',(*changes.values(),sid))
        assert not s.get('/api/assignments').json()[0]['submission']['can_delete'],changes
        assert s.delete(f'/api/submissions/{sid}').status_code==409,changes
    with a.db() as c:
        c.execute("UPDATE submissions SET status='待批改' WHERE id=?",(sid,))
        c.execute('INSERT INTO submission_history(submission_id,user_id,path,filename,snapshot,archived) VALUES(?,?,?,?,?,?)',(sid,s.get('/api/me').json()['id'],'old','old.docx','{}',a.now()))
    assert not s.get('/api/assignments').json()[0]['submission']['can_delete']
    assert s.delete(f'/api/submissions/{sid}').status_code==409

def test_student_can_delete_failed_ai_without_results_but_deadline_still_applies(monkeypatch):
    monkeypatch.setattr(a,'ai_settings',lambda:{'key':'','model':''})
    s=student('deletefailed100');assert submit(s).status_code==200
    sub=s.get('/api/assignments').json()[0];sid=sub['submission']['id']
    with a.db() as c:
        c.execute("UPDATE submissions SET status='预审需人工处理：网络失败' WHERE id=?",(sid,))
        c.execute("UPDATE assignments SET deadline='2020-01-01T00:00' WHERE id=?",(sub['id'],))
    try:
        assert s.get('/api/assignments').json()[0]['submission']['can_delete']
        assert s.delete(f'/api/submissions/{sid}').status_code==200
        assert submit(s).status_code==400
    finally:
        with a.db() as c:c.execute("UPDATE assignments SET deadline='' WHERE id=?",(sub['id'],))

def test_first_attempt_score_and_adaptive():
    s=student('score100');q=s.get('/api/practice/next').json()
    assert 'answer' not in q and 'explanation' not in q
    with a.db() as db:correct=db.execute('SELECT answer FROM questions WHERE id=?',(q['id'],)).fetchone()[0]
    r=s.post('/api/practice/answer',json={'question_id':q['id'],'choice':(correct+1)%4}).json();assert r['score']==0 and r['counted']
    r=s.post('/api/practice/answer',json={'question_id':q['id'],'choice':correct}).json();assert not r['counted']
    assert s.get('/api/progress').json()['score']==0
    assert s.get('/api/progress').json()['done']==1
    assert s.get('/api/practice/next').json()['id']!=q['id']
    assert s.post('/api/practice/answer',json={'question_id':q['id'],'choice':10}).status_code==422

def test_upload_permissions_and_publish():
    s=student('upload100');other=student('upload200');t=teacher()
    assert submit(s).status_code==200
    assert submit(s).status_code==409
    data=s.get('/api/assignments').json()[0]['submission'];sid=data['id']
    assert s.get(f'/api/submissions/{sid}/file').status_code==200
    assert other.get(f'/api/submissions/{sid}/file').status_code==403
    assert t.get(f'/api/submissions/{sid}/file').status_code==200
    assert t.post(f'/api/teacher/submissions/{sid}/grade',json={'score':85,'feedback':'草稿评语','published':False}).status_code==200
    hidden=s.get('/api/assignments').json()[0]['submission'];assert hidden['score'] is None and hidden['feedback']==''
    assert t.post(f'/api/teacher/submissions/{sid}/grade',json={'score':86,'feedback':'最终评语','published':True}).status_code==200
    visible=s.get('/api/assignments').json()[0]['submission'];assert visible['score']==86 and visible['feedback']=='最终评语'
    assert t.post(f'/api/teacher/submissions/{sid}/grade',json={'score':120,'feedback':'','published':True}).status_code==422
    assert t.delete(f'/api/teacher/submissions/{sid}').status_code==400
    assert t.post(f'/api/teacher/submissions/{sid}/return',json={'reason':'请补充实验参数和研究问题'}).status_code==200
    assert s.get(f'/api/submissions/{sid}/file').status_code==200
    assert submit(s).status_code==200
    history=t.get(f'/api/teacher/submissions/{sid}/history').json()
    assert len(history)==1 and history[0]['return_reason']=='请补充实验参数和研究问题'
    assert t.get(f"/api/teacher/history/{history[0]['id']}/file").status_code==200
    assert s.get(f"/api/teacher/history/{history[0]['id']}/file").status_code==403

def test_upload_boundaries_formats_and_deadline():
    s=student('limits100')
    assert s.post('/api/submissions',data={'assignment_id':1},files={'file':('x.html',b'<html>')}).status_code==400
    assert s.post('/api/submissions',data={'assignment_id':1},files={'file':('x.docx',b'not docx')}).status_code==400
    assert s.post('/api/submissions',data={'assignment_id':1},files={'file':('x.pdf',b'%PDF-'+b'x'*(a.LIMIT-4))}).status_code==413
    t=teacher();r=t.post('/api/assignments',json={'title':'过期作业','instructions':'要求','deadline':'2020-01-01T00:00'});assert r.status_code==200
    aid=s.get('/api/assignments').json()[0]['id']
    assert s.post('/api/submissions',data={'assignment_id':aid},files={'file':('x.docx',doc())}).status_code==400

def test_material_auth_and_reset():
    t=teacher();s=student('reset100')
    assert s.post('/api/materials',data={'chapter_id':1,'title':'课件'},files={'file':('test.pdf',b'%PDF-1.4 sample')}).status_code==403
    assert t.post('/api/materials',data={'chapter_id':1,'title':'课件'},files={'file':('test.pdf',b'%PDF-1.4 sample')}).status_code==200
    mid=s.get('/api/chapters').json()[0]['materials'][0]['id']
    assert s.get(f'/api/materials/{mid}?preview=true').headers['content-type']=='application/pdf'
    uid=s.get('/api/me').json()['id']
    assert t.post(f'/api/teacher/students/{uid}/password',json={'password':'new-password-123'}).status_code==200
    assert s.get('/api/me').status_code==401
    assert s.post('/api/login',json={'student_no':'reset100','password':'new-password-123'}).status_code==200
    assert t.delete(f'/api/materials/{mid}').status_code==200

def test_ai_workflow_mocked(monkeypatch):
    t=teacher();s=student('ai100');assert submit(s).status_code==200
    sid=next(x['id'] for x in t.get('/api/teacher/submissions').json() if x['student_no']=='ai100')
    assert t.post(f'/api/teacher/submissions/{sid}/ai').status_code==400
    assert t.post('/api/teacher/ai-settings',json={'base_url':'http://bad.example','model':'test','key':'dummy-test-key'}).status_code==400
    assert t.post('/api/teacher/ai-settings',json={'base_url':'https://example.invalid/v1','model':'test','key':'dummy-test-key','auto_review':False,'auto_return':False}).status_code==200
    assert 'dummy-test-key' not in t.get('/api/teacher/ai-settings').text
    def mock(system,payload):
        if 'document' in payload:return {'dimensions':[{'name':r['name'],'score':r['max']*.8,'comment':'测试评语'} for r in payload['rubric']],'summary':'测试初评','issues':[{'quote':'本文研究激光焊接速度与熔深的关系。','problem':'需说明范围','suggestion':'补充参数'},{'quote':'不存在的原文','problem':'测试未匹配','suggestion':'人工核查'}],'reference_note':'文献未核实'}
        return {'stem':'AI测试题','options':['甲','乙','丙','丁'],'answer':1,'explanation':'测试解析'}
    monkeypatch.setattr(a,'ai_call',mock)
    assert t.post(f'/api/teacher/submissions/{sid}/ai').status_code==200
    for _ in range(100):
        with a.db() as db:row=db.execute('SELECT * FROM submissions WHERE id=?',(sid,)).fetchone()
        if row['status']=='形式预审需教师复核':break
        time.sleep(.01)
    assert row['status']=='形式预审需教师复核'
    result=json.loads(row['ai_json']);assert result['score'] is None and result['formal_check']['decision']=='manual'
    assert row['score'] is None and row['published']==0
    q=s.post('/api/practice/ai').json();assert 'answer' not in q
    assert s.post(f"/api/practice/ai/{q['id']}/answer",json={'choice':1}).json()['score']==100
    assert s.get('/api/progress').json()['done']==0

def test_export_and_logout():
    t=teacher();s=student('export100')
    r=t.get('/api/teacher/export');assert r.status_code==200 and '平时成绩' in r.text
    assert s.get('/api/teacher/export').status_code==403
    assert s.post('/api/logout').status_code==200
    assert s.get('/api/me').status_code==401

def test_learning_profile_privacy_and_progress():
    s=student('profile100');other=student('profile200');t=teacher()
    uid=s.get('/api/me').json()['id']
    p=s.get('/api/profile').json();assert p['total_attempts']==0 and p['modules'][0]['mastery'] is None
    assert p['recommendations'][0]['level']==1
    assert s.post('/api/profile',json={'research_area':'激光制造','writing_stage':'初稿写作','goal':'补齐研究问题'}).status_code==200
    assert t.post(f'/api/teacher/students/{uid}/profile',json={'research_area':'激光制造','writing_stage':'初稿写作','goal':'补齐研究问题','teacher_note':'仅教师能看到的备注'}).status_code==200
    assert 'teacher_note' not in s.get('/api/profile').json()
    assert '仅教师能看到的备注' in t.get(f'/api/teacher/students/{uid}/profile').text
    assert other.get(f'/api/teacher/students/{uid}/profile').status_code==403
    q=s.get('/api/practice/next?chapter_id=1').json()
    with a.db() as db:correct=db.execute('SELECT answer FROM questions WHERE id=?',(q['id'],)).fetchone()[0]
    s.post('/api/practice/answer',json={'question_id':q['id'],'choice':(correct+1)%4})
    p=s.get('/api/profile').json();assert p['modules'][0]['mastery']==0 and p['weak_topics']
    s.post('/api/practice/answer',json={'question_id':q['id'],'choice':correct})
    p=s.get('/api/profile').json();assert p['modules'][0]['mastery']==32.5 and p['assessment']['score']==0
    assert p['modules'][0]['recommended_level']==2
    assert p['total_attempts']==2

def test_auto_triage_and_revision_preserves_history(monkeypatch):
    t=teacher();s=student('triage100')
    t.post('/api/teacher/ai-settings',json={'base_url':'https://example.invalid/v1','model':'test','key':'dummy','auto_review':True,'auto_return':False})
    def mock(system,payload):
        return {'dimensions':[{'name':r['name'],'score':r['max']*.4,'comment':'实验信息缺失'} for r in payload['rubric']],'summary':'需修改','issues':[],'triage':{'recommendation':'return','reasons':['尚未提供可复现的实验条件'],'required_changes':['补充设备、参数与样本量'],'evidence':['采用三组实验测试不同速度对熔深的影响，以截面金相结果作为评价依据。']}}
    monkeypatch.setattr(a,'ai_call',mock)
    r=submit(s);assert r.status_code==200 and r.json()['auto_review']
    sid=s.get('/api/assignments').json()[-1]['submission']['id']
    for _ in range(100):
        rows=t.get('/api/teacher/submissions').json();row=next(x for x in rows if x['id']==sid)
        if row['status']=='形式预审需教师复核':break
        time.sleep(.01)
    assert row['ai']['formal_check']['decision']=='manual' and row['returned']==0
    assert s.get('/api/profile').json()['writing_reviews']==[]
    visible=next(x for x in s.get('/api/assignments').json() if x['id']==1)['submission']
    assert visible['status']=='形式预审需教师复核' and visible['score'] is None
    assert t.post(f'/api/teacher/submissions/{sid}/return',json={'reason':'请补充设备、参数与样本量后重交'}).status_code==200
    visible=next(x for x in s.get('/api/assignments').json() if x['id']==1)['submission']
    assert visible['returned']==1 and '样本量' in visible['return_reason']
    assert t.post(f'/api/teacher/submissions/{sid}/grade',json={'score':50,'feedback':'','published':True}).status_code==409
    assert submit(s).status_code==200
    for _ in range(100):
        with a.db() as db:status=db.execute('SELECT status FROM submissions WHERE id=?',(sid,)).fetchone()[0]
        if status=='形式预审需教师复核':break
        time.sleep(.01)
    history=t.get(f'/api/teacher/submissions/{sid}/history').json();assert len(history)==1 and history[0]['ai']['formal_check']['decision']=='manual'
    t.post('/api/teacher/ai-settings',json={'base_url':'https://example.invalid/v1','model':'test','key':'','auto_review':False,'auto_return':False})

def test_triage_requires_verified_evidence(monkeypatch):
    t=teacher();s=student('evidence100');assert submit(s).status_code==200
    sid=next(x['id'] for x in t.get('/api/teacher/submissions').json() if x['student_no']=='evidence100')
    def mock(system,payload):return {'dimensions':[{'name':r['name'],'score':0,'comment':'低分测试'} for r in payload['rubric']],'triage':{'recommendation':'return','reasons':['质量低'],'required_changes':['重新写作'],'evidence':['不存在的原文证据']}}
    monkeypatch.setattr(a,'ai_call',mock)
    a.run_review(sid)
    row=next(x for x in t.get('/api/teacher/submissions').json() if x['id']==sid)
    assert row['ai']['score'] is None and row['ai']['triage']['recommendation']=='review'

@pytest.mark.parametrize('archive_name',['writing-class-v1.zip','writing-class-v1.1.zip'])
def test_upgrade_from_v1_retains_data(tmp_path,archive_name):
    import subprocess,sys,zipfile
    from pathlib import Path
    archive=Path(__file__).parent.parent/archive_name
    if not archive.exists():
        import pytest
        pytest.skip('旧版发布包不在当前测试目录，跳过迁移夹具')
    legacy=tmp_path/'legacy';legacy.mkdir()
    with zipfile.ZipFile(archive) as z:z.extractall(legacy)
    code=legacy/('writing-class-v1.1' if 'v1.1' in archive_name else 'writing-class');data=tmp_path/'legacy-data'
    env={**os.environ,'DATA_DIR':str(data),'PYTHONPATH':str(code)}
    script="""import app
with app.db() as c:
 c.execute(\"INSERT INTO users(id,name,student_no,password,role,created) VALUES(101,'旧学生','legacy101','original-password-hash','student','2026-10-01')\")
 c.execute(\"INSERT INTO attempts(user_id,question_id,choice,score,created) VALUES(101,1,2,100,'2026-10-01')\")
 c.execute(\"INSERT INTO submissions(assignment_id,user_id,path,filename,created,score,feedback,published) VALUES(1,101,'/old/file.docx','旧稿.docx','2026-10-01',88,'原评语',1)\")
"""
    subprocess.run([sys.executable,'-c',script],cwd=code,env=env,check=True,capture_output=True)
    (code/'app.py').write_text(Path(a.__file__).read_text())
    (code/'review_pipeline.py').write_text((Path(a.__file__).parent/'review_pipeline.py').read_text())
    import shutil
    for module in ('english_course.py','english_speech.py','paper_queue.py','writing_workflow.py','writing_slides.py','adaptive_practice.py','practice_bank.py'):
        shutil.copy2(Path(a.__file__).parent/module,code/module)
    shutil.copytree(Path(a.__file__).parent/'vendor',code/'vendor',dirs_exist_ok=True)
    shutil.copytree(Path(a.__file__).parent/'question_bank',code/'question_bank',dirs_exist_ok=True)
    verify="""import app,json
with app.db() as c:
 u=c.execute('SELECT * FROM users WHERE id=101').fetchone()
 s=c.execute('SELECT * FROM submissions WHERE user_id=101').fetchone()
 assert u['password']=='original-password-hash'
 assert s['score']==88 and s['feedback']=='原评语' and s['published']==1 and s['returned']==0
 assert s['path']=='/old/file.docx'
 p=app.learner_profile(101,c)
 assert p['assessment']['done']==1 and p['total_attempts']==1
 print('migration-ok')
"""
    result=subprocess.run([sys.executable,'-c',verify],cwd=code,env=env,check=True,capture_output=True,text=True)
    assert 'migration-ok' in result.stdout


import pytest
from contextlib import contextmanager
@pytest.fixture(autouse=True)
def docx_render_fixture(monkeypatch,tmp_path):
    with a.db() as c:c.execute('DELETE FROM throttle')
    original=a.prepare_document
    @contextmanager
    def wrapped(path):
        if str(path).endswith('.docx'):
            from docx import Document
            from PIL import Image
            image=tmp_path/'page.jpg';Image.new('RGB',(800,1000),'white').save(image)
            text='\n'.join(p.text for p in Document(path).paragraphs)
            yield {'text':text,'pages':[{'page_no':1,'text':text,'image_path':str(image)}],'page_count':1,'complete':True,'warnings':[],'word_rendered':True}
        else:
            with original(path) as d:yield d
    monkeypatch.setattr(a,'prepare_document',wrapped)

def formal_response(quote,decision='revise',count=5,confidence=.96):
    return {'summary':'存在多处明确形式问题','formal_check':{'decision':decision,'confidence':confidence,'readable':True,'issues':[{'page':1,'location':f'第{i+1}段','category':'expression','severity':'major' if i<2 else 'minor','problem':f'第{i+1}处术语表述存在明确不一致','change':f'请统一第{i+1}处术语并按课程要求修正','quote':quote,'evidence_type':'text','evidence':''} for i in range(count)]}}

def test_v12_auto_return_checklist_and_appeal(monkeypatch):
    t=teacher();s=student('formal100');other=student('formal200')
    t.post('/api/teacher/ai-settings',json={'base_url':'https://api.deepseek.com','model':'deepseek-flash','key':'mock-key','auto_review':True,'auto_return':True})
    monkeypatch.setattr(a,'ai_call',lambda system,payload:formal_response('本文研究激光焊接速度与熔深的关系。'))
    assert submit(s).status_code==200
    for _ in range(150):
        sub=next(x for x in s.get('/api/assignments').json() if x['id']==1)['submission']
        if sub['returned']:break
        time.sleep(.01)
    assert sub['returned']==1 and '第1页' in sub['return_reason'] and '修改要求' in sub['return_reason']
    sid=sub['id'];assert sub['score'] is None
    with a.db() as db:assert db.execute('SELECT score,published FROM submissions WHERE id=?',(sid,)).fetchone()['score'] is None
    assert other.post(f'/api/submissions/{sid}/appeal',json={'reason':'请求人工重新查看论文'}).status_code==404
    assert s.post(f'/api/submissions/{sid}/appeal',json={'reason':'页码中的术语实际一致，请教师核对'}).status_code==200
    row=next(x for x in t.get('/api/teacher/submissions').json() if x['id']==sid)
    assert row['returned']==0 and row['status']=='学生申请教师复核' and row['appeal_reason']
    assert t.post(f'/api/teacher/submissions/{sid}/grade',json={'score':80,'feedback':'经人工复核','published':True}).status_code==200
    t.post('/api/teacher/ai-settings',json={'base_url':'https://example.invalid/v1','model':'test','key':'dummy','auto_review':False,'auto_return':False})

def test_formal_decision_threshold_and_uncertainty():
    from review_pipeline import normalize_formal
    quote='原文内容用于测试引用验证。'
    document={'pages':[{'page_no':1,'text':quote}],'page_count':1,'complete':True,'warnings':[]}
    f=normalize_formal(formal_response(quote)['formal_check'],document)
    assert f['automatic_eligible'] and f['located_count']==5
    assert not normalize_formal(formal_response(quote,count=4)['formal_check'],document)['automatic_eligible']
    assert not normalize_formal(formal_response(quote,confidence=.89)['formal_check'],document)['automatic_eligible']
    assert not normalize_formal(formal_response('编造的不存在原文')['formal_check'],document)['automatic_eligible']
    raw=formal_response(quote)['formal_check'];raw['readable']=False
    assert normalize_formal(raw,document)['decision']=='manual'
    raw=formal_response(quote)['formal_check'];raw['issues']=[raw['issues'][0]]*8
    assert normalize_formal(raw,document)['located_count']==1
    raw=formal_response(quote)['formal_check'];raw['issues'][0]['page']=99
    assert not normalize_formal(raw,document)['automatic_eligible']

def test_real_pdf_and_docx_rendering(tmp_path):
    import pymupdf, shutil
    from review_pipeline import prepare_document
    pdf=tmp_path/'paper.pdf';d=pymupdf.open();p=d.new_page();p.insert_text((70,70),'Academic paper: title, abstract and figure 1.');d.save(pdf);d.close()
    with prepare_document(pdf) as result:
        assert result['page_count']==1 and 'Academic' in result['text']
        image=result['pages'][0]['image_path'];assert os.path.getsize(image)>100
    assert not os.path.exists(image)
    if shutil.which('soffice') or shutil.which('libreoffice'):
        word=tmp_path/'paper.docx';word.write_bytes(doc())
        with prepare_document(word) as result:
            assert result['page_count']>=1 and result['word_rendered'] and result['warnings']
            assert '激光焊接' in result['text']

def test_long_pdf_routes_to_manual(tmp_path):
    import pymupdf
    from review_pipeline import prepare_document
    path=tmp_path/'long.pdf';d=pymupdf.open()
    for _ in range(21):d.new_page()
    d.save(path);d.close()
    with pytest.raises(ValueError,match='人工'):
        with prepare_document(path):pass

def test_multimodal_payload_usage_and_local_cache(monkeypatch,tmp_path):
    from PIL import Image
    t=teacher()
    t.post('/api/teacher/ai-settings',json={'base_url':'https://api.deepseek.com','model':'deepseek-flash','key':'mock-key','auto_review':False,'auto_return':False})
    image=tmp_path/'figure.jpg';Image.new('RGB',(40,40),'white').save(image)
    requests=[];original=a.httpx.Client
    def handler(request):
        x=json.loads(request.content);requests.append(x)
        assert x['thinking']=={'type':'disabled'}
        assert x['messages'][1]['content'][1]['type']=='image_url'
        assert x['messages'][1]['content'][1]['image_url']['url'].startswith('data:image/jpeg;base64,')
        return a.httpx.Response(200,json={'choices':[{'finish_reason':'stop','message':{'content':'{"formal_check":{"decision":"pass"}}'}}],'usage':{'prompt_tokens':1200,'completion_tokens':100,'prompt_cache_hit_tokens':200}})
    monkeypatch.setattr(a.httpx,'Client',lambda **kwargs:original(transport=a.httpx.MockTransport(handler)))
    payload={'document':'unique-multimodal-test','_images':[str(image)],'_usage_context':{'user_id':99901,'purpose':'测试图文'}}
    one=a.ai_call('JSON test',payload);two=a.ai_call('JSON test',payload)
    assert one==two and len(requests)==1
    rows=t.get('/api/teacher/ai-usage').json()['records']
    assert rows[0]['local_cache']==1 and rows[0]['input_tokens']==0
    assert rows[1]['input_tokens']==1200 and rows[1]['output_tokens']==100

def test_five_mib_boundary():
    s=student('five100')
    assert s.post('/api/submissions',data={'assignment_id':1},files={'file':('five.pdf',b'%PDF-'+b'x'*(a.LIMIT-5))}).status_code==200
    s2=student('five200')
    assert s2.post('/api/submissions',data={'assignment_id':1},files={'file':('large.pdf',b'%PDF-'+b'x'*(a.LIMIT-4))}).status_code==413

def test_four_course_catalog_and_placeholders():
    anonymous=TestClient(a.app,headers=H)
    assert anonymous.get('/api/courses').status_code==401
    s=student('courses100')
    rows=s.get('/api/courses').json()
    assert [x['title'] for x in rows]==['学术论文写作','机械工程专业英语','智能装备基础','智能制造课程设计']
    assert {x['id'] for x in rows if x['state']=='planned'}=={'equipment'}
    assert s.get('/api/chapters').status_code==200
    assert s.get('/api/assignments').status_code==200

def test_gomoku_local_assets_authentication_and_paths():
    import hashlib,re
    anonymous=TestClient(a.app,headers=H)
    assert anonymous.get('/courses/design/gomoku/').status_code==401
    assert anonymous.get('/courses/design/gomoku/gomoku_uint8_qop_u8u8.onnx').status_code==401
    s=student('gomoku100')
    html=s.get('/courses/design/gomoku/')
    assert html.status_code==200 and 'text/html' in html.headers['content-type']
    assert '智能制造案例' in html.text and 'ai-worker.js' in html.text
    js=s.get('/courses/design/gomoku/ai-worker.js?v=safe-vcf-tested-20260911-1')
    assert js.status_code==200 and 'application/javascript' in js.headers['content-type']
    model=s.get('/courses/design/gomoku/gomoku_uint8_qop_u8u8.onnx?v=safe-vcf-tested-20260911-1')
    assert model.status_code==200 and len(model.content)==150730
    assert hashlib.sha1(b'blob '+str(len(model.content)).encode()+b'\0'+model.content).hexdigest()=='43815dcdfafdf2a41528d00058b1c7af4be33b09'
    assert s.get('/courses/design/gomoku/%2e%2e%2fapp.py').status_code==404
    assert s.get('/courses/design/gomoku/not-there.wasm').status_code==404
    assert s.get('/courses/design/gomoku/index.html').status_code==200
    assert s.get('/courses/design/gomoku/gomoku_single.onnx').status_code==200
    assert s.post('/api/logout').status_code==200
    assert s.get('/courses/design/gomoku/ai-worker.js').status_code==401
